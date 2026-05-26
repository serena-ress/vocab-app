"""从 Excel / PDF 文件解析单词列表。"""

from __future__ import annotations

import re
from pathlib import Path

WordRow = tuple[str, str, str]

HEADER_ALIASES: dict[str, set[str]] = {
    "word": {"word", "单词", "英文", "english", "词汇", "词"},
    "definition": {
        "definition",
        "释义",
        "中文",
        "意思",
        "meaning",
        "翻译",
        "解释",
    },
    "example": {"example", "例句", "sentence", "例", "句子"},
}

_LINE_PATTERNS = [
    re.compile(r"^(.+?)\t(.+?)(?:\t(.+))?$"),
    re.compile(r"^(.+?)\|(.+?)(?:\|(.+))?$"),
    re.compile(r"^(.+?)[,，;；](.+?)(?:[,，;；](.+))?$"),
    re.compile(r"^(.+?)\s+[-–—]\s+(.+)$"),
]


def _cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _match_header(cell: str) -> str | None:
    key = cell.strip().lower()
    for field, aliases in HEADER_ALIASES.items():
        if key in aliases or key == field:
            return field
    return None


def _column_map(header_row: tuple) -> dict[str, int] | None:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        field = _match_header(_cell_str(cell))
        if field and field not in mapping:
            mapping[field] = idx
    if "word" in mapping and "definition" in mapping:
        return mapping
    return None


def parse_text_lines(text: str) -> list[WordRow]:
    """解析纯文本行：支持 Tab/竖线/逗号/破折号分隔。"""
    rows: list[WordRow] = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        line = re.sub(r"^\d+[\.\)、]\s*", "", line).strip()
        if not line:
            continue

        parsed: WordRow | None = None
        for pattern in _LINE_PATTERNS:
            m = pattern.match(line)
            if m:
                groups = m.groups()
                word = groups[0].strip()
                definition = groups[1].strip()
                example = (groups[2] or "").strip() if len(groups) > 2 else ""
                parsed = (word, definition, example)
                break

        if parsed and parsed[0] and parsed[1]:
            rows.append(parsed)
    return rows


def import_from_excel(path: Path) -> list[WordRow]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("请先安装 openpyxl：pip install openpyxl") from exc

    if path.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        raise ValueError("仅支持 .xlsx 格式，请将 .xls 另存为 .xlsx 后导入。")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        raw_rows = [
            tuple(_cell_str(c) for c in row)
            for row in ws.iter_rows(values_only=True)
            if any(_cell_str(c) for c in row)
        ]
    finally:
        wb.close()

    if not raw_rows:
        return []

    col_map = _column_map(raw_rows[0])
    data_rows = raw_rows[1:] if col_map else raw_rows
    if not col_map:
        col_map = {"word": 0, "definition": 1, "example": 2}

    rows: list[WordRow] = []
    for row in data_rows:
        word = row[col_map["word"]] if col_map["word"] < len(row) else ""
        definition = row[col_map["definition"]] if col_map["definition"] < len(row) else ""
        ex_idx = col_map.get("example")
        example = row[ex_idx] if ex_idx is not None and ex_idx < len(row) else ""
        if word and definition:
            rows.append((word, definition, example))
    return rows


def import_from_pdf(path: Path) -> list[WordRow]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError("请先安装 pypdf：pip install pypdf") from exc

    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            parts.append(text)
    if not parts:
        raise ValueError("未能从 PDF 中提取文字，请使用可选中文字的 PDF。")

    return parse_text_lines("\n".join(parts))
